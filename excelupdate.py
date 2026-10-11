import sqlite3
import time
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

DB_PATH = Path("core.db")
EXCEL_PATH = Path("data") / "MineGuard_Telemetry.xlsx"
REFRESH_SECONDS = 30


def export_to_excel():
    EXCEL_PATH.parent.mkdir(parents=True, exist_ok=True)

    with sqlite3.connect(DB_PATH, timeout=10) as conn:
        wb = (
            load_workbook(EXCEL_PATH)
            if EXCEL_PATH.exists()
            else Workbook()
        )

        if "Summary" not in wb.sheetnames:
            ws = wb.active
            ws.title = "Summary"
        else:
            ws = wb["Summary"]

        ws.delete_rows(1, ws.max_row)
        ws.append(["MineGuard Local Data Export"])
        ws.append([
            "Last refreshed",
            datetime.now().astimezone().isoformat()
        ])
        ws.append(["Database", str(DB_PATH.resolve())])

        for table, sheet_name in [
            ("telemetry", "Telemetry"),
            ("alerts", "Alerts"),
        ]:
            if sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                sheet.delete_rows(1, sheet.max_row)
            else:
                sheet = wb.create_sheet(sheet_name)

            cursor = conn.execute(f'SELECT * FROM "{table}"')
            sheet.append([col[0] for col in cursor.description])

            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(
                    "solid", fgColor="1F4E78"
                )
                cell.alignment = Alignment(wrap_text=True)

            count = 0
            for row in cursor:
                sheet.append(list(row))
                count += 1

            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions

            for cells in sheet.columns:
                letter = cells[0].column_letter
                length = max(
                    (len(str(c.value)) for c in cells
                     if c.value is not None),
                    default=10,
                )
                sheet.column_dimensions[letter].width = min(
                    max(length + 2, 12), 40
                )

            ws.append([f"{table} records", count])

        wb.save(EXCEL_PATH)
        wb.close()

    print(f"Excel refreshed: {EXCEL_PATH}")


def main():
    if not DB_PATH.exists():
        raise SystemExit(f"Database not found: {DB_PATH}")

    while True:
        try:
            export_to_excel()
        except PermissionError:
            print(
                "Excel file is open or locked. "
                "Close it and the next refresh will retry."
            )
        except Exception as exc:
            print(f"Export error: {exc}")

        time.sleep(REFRESH_SECONDS)


if __name__ == "__main__":
    main()
